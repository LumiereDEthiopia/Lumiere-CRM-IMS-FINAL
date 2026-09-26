import openpyxl, json, sys

path = r"c:\Users\leult\OneDrive\Desktop\ELMAS\course for me\advanced stratagy\lumiere-perfumes-2026-09-12.xlsx"
wb = openpyxl.load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print(f"\n===== SHEET: {ws.title}  dims={ws.dimensions} rows={ws.max_row} cols={ws.max_column} =====")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60), values_only=False):
        vals = []
        for c in row:
            v = c.value
            if v is None:
                v = ""
            else:
                v = str(v).replace("\n", "\\n")
            vals.append(v)
        # trim trailing empties
        while vals and vals[-1] == "":
            vals.pop()
        if vals:
            print(" | ".join(vals))